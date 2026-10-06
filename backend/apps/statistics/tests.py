import io

from django.contrib.auth import get_user_model
from django.urls import reverse
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APITestCase

from apps.academics.models import AcademicYear, Cohort, Department, Field, Semester
from apps.topics.models import Topic


class StatisticsOverviewTests(APITestCase):
    def setUp(self):
        User = get_user_model()
        self.admin = User.objects.create_user(username="stats_admin", role="admin")
        self.head = User.objects.create_user(username="stats_head", role="department_head")
        self.teacher_one = User.objects.create_user(
            username="teacher_one", first_name="An", last_name="Nguyễn", role="teacher"
        )
        self.teacher_two = User.objects.create_user(
            username="teacher_two", first_name="An", last_name="Nguyễn", role="teacher"
        )
        self.department = Department.objects.create(name="Phần mềm", code="PM")
        self.other_department = Department.objects.create(name="Mạng máy tính", code="MMT")
        self.cohort = Cohort.objects.create(name="K80", start_year=2040, end_year=2044)
        self.other_cohort = Cohort.objects.create(name="K81", start_year=2041, end_year=2045)
        self.year = AcademicYear.objects.create(name="2040-2041", is_current=True)
        self.other_year = AcademicYear.objects.create(name="2041-2042")
        self.semester = Semester.objects.create(
            name="HK1", academic_year=self.year, start_date="2040-09-01", end_date="2041-01-31"
        )
        self.other_semester = Semester.objects.create(
            name="HK1", academic_year=self.other_year, start_date="2041-09-01", end_date="2042-01-31"
        )
        self.field = Field.objects.create(name="Dữ liệu")
        Topic.objects.create(
            title="Đề tài 1", department=self.department, field=self.field,
            cohort=self.cohort, academic_year=self.year, semester=self.semester,
            proposed_by=self.teacher_one, status=Topic.Status.PENDING,
        )
        Topic.objects.create(
            title="Đề tài 2", department=self.department, field=self.field,
            cohort=self.cohort, academic_year=self.year, semester=self.semester,
            proposed_by=self.teacher_two, status=Topic.Status.APPROVED,
        )
        Topic.objects.create(
            title="Đề tài 3", department=self.other_department,
            cohort=self.other_cohort, academic_year=self.other_year, semester=self.other_semester,
            proposed_by=self.teacher_one, status=Topic.Status.REJECTED,
        )
        self.url = reverse("statistics-overview")

    def test_admin_sees_all_aggregation_dimensions(self):
        self.client.force_authenticate(self.admin)

        response = self.client.get(self.url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["total"], 3)
        self.assertEqual(len(response.data["by_academic_year"]), 2)
        self.assertEqual(len(response.data["by_cohort"]), 2)
        self.assertEqual(len(response.data["by_teacher"]), 2)
        self.assertEqual(len(response.data["by_field"]), 2)
        self.assertEqual(len(response.data["by_status"]), 3)
        self.assertEqual(
            {row["proposed_by__username"] for row in response.data["by_teacher"]},
            {"teacher_one", "teacher_two"},
        )

    def test_filters_update_statistics_and_empty_result_is_zero(self):
        self.client.force_authenticate(self.admin)

        filtered = self.client.get(self.url, {
            "academic_year": self.year.pk,
            "proposed_by": self.teacher_one.pk,
        })
        empty = self.client.get(self.url, {"status": Topic.Status.RENAME_REQUESTED})

        self.assertEqual(filtered.data["total"], 1)
        self.assertEqual(filtered.data["by_teacher"][0]["proposed_by__username"], "teacher_one")
        self.assertEqual(empty.status_code, status.HTTP_200_OK)
        self.assertEqual(empty.data["total"], 0)
        self.assertEqual(empty.data["by_status"], [])

    def test_statistics_are_admin_only(self):
        for user in (self.head, self.teacher_one):
            with self.subTest(role=user.role):
                self.client.force_authenticate(user)
                self.assertEqual(self.client.get(self.url).status_code, status.HTTP_403_FORBIDDEN)

    def test_authenticated_excel_export_honors_filters(self):
        self.client.force_authenticate(self.admin)

        response = self.client.get(reverse("statistics-export"), {
            "export_format": "excel",
            "department": self.other_department.pk,
            "semester": self.other_semester.pk,
            "search": "Đề tài 3",
        })

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = load_workbook(io.BytesIO(response.content))
        sheet = workbook.active
        self.assertEqual(sheet["A1"].value, "BÁO CÁO DANH SÁCH ĐỀ TÀI")
        self.assertEqual(sheet["A4"].value, "ID")
        self.assertEqual(sheet["B5"].value, "Đề tài 3")
        self.assertEqual(sheet.max_row, 5)
        self.assertIn("Học kỳ: HK1", sheet["A3"].value)

    def test_empty_excel_export_is_valid_and_contains_headers(self):
        self.client.force_authenticate(self.admin)

        response = self.client.get(reverse("statistics-export"), {
            "export_format": "excel",
            "search": "không tồn tại",
        })

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        sheet = load_workbook(io.BytesIO(response.content)).active
        self.assertEqual(sheet["B4"].value, "Tên đề tài")
        self.assertIn("Không có đề tài", sheet["A5"].value)

    def test_pdf_export_is_valid_and_admin_only(self):
        url = reverse("statistics-export")
        self.client.force_authenticate(self.admin)
        response = self.client.get(url, {"export_format": "pdf", "status": Topic.Status.PENDING})

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF"))

        empty_response = self.client.get(url, {"export_format": "pdf", "search": "không tồn tại"})
        self.assertEqual(empty_response.status_code, status.HTTP_200_OK)
        self.assertTrue(empty_response.content.startswith(b"%PDF"))

        self.client.force_authenticate(self.head)
        self.assertEqual(
            self.client.get(url, {"export_format": "excel"}).status_code,
            status.HTTP_403_FORBIDDEN,
        )
