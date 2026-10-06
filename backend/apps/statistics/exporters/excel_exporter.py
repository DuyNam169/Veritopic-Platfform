"""Xuất danh sách đề tài ra file Excel (.xlsx) bằng openpyxl."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def _safe_cell(value):
    """Prevent values supplied by users from being interpreted as Excel formulas."""
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def export_topics_to_excel(topics, *, generated_at="", filter_summary="Tất cả đề tài") -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Danh sách đề tài"

    headers = ["ID", "Tên đề tài", "Bộ môn", "Lĩnh vực", "Khóa", "Năm học", "Học kỳ", "GVHD", "Trạng thái"]
    ws.merge_cells("A1:I1")
    ws["A1"] = "BÁO CÁO DANH SÁCH ĐỀ TÀI"
    ws["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="1E3A8A")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30
    ws.merge_cells("A2:I2")
    ws["A2"] = f"Thời điểm xuất: {generated_at}"
    ws.merge_cells("A3:I3")
    ws["A3"] = f"Phạm vi dữ liệu: {filter_summary}"
    ws["A3"].alignment = Alignment(wrap_text=True)

    ws.append(headers)
    header_row = 4
    thin = Side(style="thin", color="CBD5E1")
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

    row_count = 0
    for t in topics:
        row_count += 1
        ws.append(
            [
                t.id,
                _safe_cell(t.title),
                t.department.name if t.department else "",
                t.field.name if t.field else "",
                t.cohort.name if t.cohort else "",
                t.academic_year.name if t.academic_year else "",
                t.semester.name if t.semester else "",
                _safe_cell((t.proposed_by.get_full_name() or t.proposed_by.username) if t.proposed_by else ""),
                t.get_status_display(),
            ]
        )

    if row_count == 0:
        ws.merge_cells("A5:I5")
        ws["A5"] = "Không có đề tài phù hợp với điều kiện lọc."
        ws["A5"].alignment = Alignment(horizontal="center")
        ws["A5"].font = Font(italic=True, color="64748B")

    for row in ws.iter_rows(min_row=5):
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    widths = [10, 45, 24, 24, 14, 16, 16, 28, 20]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:I{max(ws.max_row, 4)}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:4"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
